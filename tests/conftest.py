"""Shared fixtures: a small workbook laid out like the ONS healthy life expectancy file.

The sample has the same sheet names and header rows the script expects, but only a few
rows, so tests run in milliseconds and the expected results can be worked out by hand.
"""

from pathlib import Path

import pytest
from openpyxl import Workbook

TABLE_COLUMNS = [
    "Period",
    "Country",
    "Area type",
    "Area code",
    "Area name",
    "Sex",
    "Sex code",
    "Age group",
    "Age code",
    "HLE",
    "LCI",
    "UCI",
    "Proportion (%)",
]

# Three areas: the tests keep two of them and expect the third to be filtered out
AREAS = [
    ("E08000011", "Knowsley"),
    ("E06000006", "Halton"),
    ("E06000001", "Hartlepool"),
]
PERIODS = ["2011 to 2013", "2022 to 2024"]
SEXES = [("Male", 1), ("Female", 2)]
AGE_GROUPS = ["<1", "01 to 04"]

# Rows above the header row on sheets 1 and 3 (the script uses header_row = 6)
PREAMBLE_ROWS = 6


def _add_text_sheet(workbook, title, prefix, line_count):
    """Add a sheet of single-cell lines like "notes line 0", "notes line 1"..."""
    sheet = workbook.create_sheet(title)
    for i in range(line_count):
        sheet.append([f"{prefix} line {i}"])


def _add_table_sheet(workbook, title, columns, rows):
    """Add a sheet with preamble lines, then a header row, then the data rows."""
    sheet = workbook.create_sheet(title)
    for i in range(PREAMBLE_ROWS):
        sheet.append([f"sheet{title} line {i}"])
    sheet.append(columns)
    for row in rows:
        sheet.append(row)


def _sheet_one_rows():
    """One row per area, period, sex and age group, with easy-to-check numbers."""
    rows = []
    n = 0
    for code, name in AREAS:
        for period in PERIODS:
            for sex, sex_code in SEXES:
                for age_code, age_group in enumerate(AGE_GROUPS, start=1):
                    hle = 50.0 + n
                    rows.append(
                        [
                            period,
                            "England",
                            "Local Areas",
                            code,
                            name,
                            sex,
                            sex_code,
                            age_group,
                            age_code,
                            hle,
                            hle - 2,
                            hle + 2,
                            70 + n % 5,
                        ]
                    )
                    n += 1
    return rows


def build_sample_workbook(path: Path) -> None:
    """Write a workbook with the same sheets as the ONS file, but tiny."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    _add_text_sheet(workbook, "Cover_sheet", "cover", 5)
    _add_text_sheet(workbook, "Contents", "contents", 4)
    _add_text_sheet(workbook, "Notes", "notes", 6)
    _add_table_sheet(workbook, "1", TABLE_COLUMNS, _sheet_one_rows())
    _add_text_sheet(workbook, "2", "sheet2", 10)  # the pivot table sheet
    _add_table_sheet(
        workbook,
        "3",
        ["Area code", "Area name", "Change"],
        [[code, name, 1.5] for code, name in AREAS],
    )
    workbook.save(path)


@pytest.fixture
def sample_workbook(tmp_path: Path) -> Path:
    """Path to a freshly built sample workbook, in its own folder."""
    directory = tmp_path / "source"
    directory.mkdir()
    path = directory / "sample.xlsx"
    build_sample_workbook(path)
    return path
